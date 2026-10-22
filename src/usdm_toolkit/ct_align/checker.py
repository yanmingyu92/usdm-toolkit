"""Core consistency checks between dataStructure.yml and USDM_CT.xlsx.

The checker compares the USDM model export (``dataStructure.yml``) against the
CDISC controlled terminology workbook (``USDM_CT.xlsx``, sheet
"DDF Entities&Attributes").

Rules implemented (rule ids are stable and used in reports):

- ``CLASS_NOT_IN_CT``: a class in dataStructure.yml has no corresponding entity
  row in USDM_CT.xlsx (e.g. API-only classes such as ExtensionAttribute /
  ExtensionClass, cf. cdisc-org/DDF-RA#718).
- ``CLASS_CCODE_MISMATCH`` / ``CLASS_TERM_MISMATCH`` / ``CLASS_DEFINITION_MISMATCH``:
  the class-level NCI C-code / preferred term / definition differs between the
  YAML and the CT workbook.
- ``ATTR_NOT_IN_CT``: a YAML attribute carrying an NCI C-code has no matching
  (entity, attribute) row in the CT workbook.
- ``ATTR_CCODE_MISMATCH`` / ``ATTR_TERM_MISMATCH`` / ``ATTR_DEFINITION_MISMATCH``:
  the attribute-level C-code / preferred term / definition differs. For C-code
  mismatches the checker also reports which CT row actually owns the YAML's
  C-code (the signature described in cdisc-org/DDF-RA#716, where attributes on
  abstract classes carry a concrete subclass's C-code).
- ``CCODE_NOT_IN_CT``: a C-code referenced from the YAML does not occur anywhere
  in the CT workbook.
- ``CT_ATTR_NOT_IN_YAML``: a CT attribute row whose counterpart in the YAML is
  missing or carries no C-code (coverage gap, informational).
- ``CLASSIFICATION_WITHOUT_VALUE_LIST``: CT rows whose definition follows the
  "A characterization or classification of ..." pattern but that declare no
  value list (Has Value List = N), cf. cdisc-org/DDF-RA#717.
- ``VALUE_LIST_TYPE_MISMATCH``: CT declares a value list (Has Value List = Y...)
  but the YAML attribute is not typed as Code/AliasCode, or vice versa.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any

import yaml

CT_SHEET_ENTITIES = "DDF Entities&Attributes"
CLASSIFICATION_PREFIX = "A characterization or classification of"

CODE_TYPES = {"#/Code", "#/AliasCode"}


@dataclass
class CtRow:
    """One row of the 'DDF Entities&Attributes' sheet."""

    row_number: int  # value of the 'Row #' column
    entity: str
    role: str  # Entity | Attribute | Relationship | Complex Datatype Relationship
    name: str  # Logical Data Model Name (attribute name; same as entity for entities)
    ccode: str | None
    preferred_term: str | None
    definition: str | None
    has_value_list: str | None


@dataclass
class CtData:
    entities: dict[str, CtRow] = field(default_factory=dict)
    attributes: dict[tuple[str, str], CtRow] = field(default_factory=dict)
    code_index: dict[str, list[str]] = field(default_factory=dict)

    def code_owners(self, ccode: str) -> list[str]:
        return self.code_index.get(ccode, [])


@dataclass
class Finding:
    rule: str
    severity: str  # "error" | "warning" | "info"
    location: str  # "Class" or "Class.attribute"
    ccode_yaml: str | None
    ccode_ct: str | None
    yaml_value: str | None
    ct_value: str | None
    message: str


def load_yaml_model(path: str) -> dict[str, Any]:
    with open(path, encoding="utf-8") as fh:
        doc = yaml.safe_load(fh)
    if not isinstance(doc, dict):
        raise ValueError(f"{path}: expected a mapping of class definitions")
    return doc


def load_ct(path: str) -> CtData:
    """Parse USDM_CT.xlsx ('DDF Entities&Attributes' sheet)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    if CT_SHEET_ENTITIES not in wb.sheetnames:
        raise ValueError(f"{path}: sheet {CT_SHEET_ENTITIES!r} not found")
    ws = wb[CT_SHEET_ENTITIES]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    col = {name: idx for idx, name in enumerate(header)}

    def get(row: tuple, name: str) -> str | None:
        value = row[col[name]]
        if value is None:
            return None
        return str(value).strip()

    data = CtData()
    for row in rows[1:]:
        if row[col["Entity Name"]] is None:
            continue
        ct_row = CtRow(
            row_number=int(get(row, "Row #") or 0),
            entity=get(row, "Entity Name") or "",
            role=get(row, "Role") or "",
            name=get(row, "Logical Data Model Name") or "",
            ccode=get(row, "NCI C-code"),
            preferred_term=get(row, "CT Item Preferred Name"),
            definition=get(row, "Definition"),
            has_value_list=get(row, "Has Value List"),
        )
        location = ct_row.entity if ct_row.role == "Entity" else f"{ct_row.entity}.{ct_row.name}"
        if ct_row.role == "Entity":
            data.entities[ct_row.entity] = ct_row
        else:
            data.attributes[(ct_row.entity, ct_row.name)] = ct_row
        if ct_row.ccode:
            data.code_index.setdefault(ct_row.ccode, []).append(location)
    return data


def _type_refs(attr_body: dict) -> set[str]:
    return {t.get("$ref", "") for t in attr_body.get("Type") or [] if isinstance(t, dict)}


def check(model: dict[str, Any], ct: CtData) -> list[Finding]:
    findings: list[Finding] = []

    def add(rule, severity, location, ccode_yaml, ccode_ct, yaml_value, ct_value, message):
        findings.append(Finding(rule, severity, location, ccode_yaml, ccode_ct,
                                yaml_value, ct_value, message))

    # --- class level -------------------------------------------------------
    for cls, body in model.items():
        if not isinstance(body, dict):
            continue
        ct_entity = ct.entities.get(cls)
        yaml_ccode = body.get("NCI C-Code")
        if ct_entity is None:
            add("CLASS_NOT_IN_CT", "warning", cls, yaml_ccode, None,
                body.get("Preferred Term"), None,
                f"Class {cls!r} has no entity row in USDM_CT.xlsx "
                f"(C-code in YAML: {yaml_ccode or 'none'})")
            continue
        if yaml_ccode and ct_entity.ccode and yaml_ccode != ct_entity.ccode:
            add("CLASS_CCODE_MISMATCH", "error", cls, yaml_ccode, ct_entity.ccode,
                body.get("Preferred Term"), ct_entity.preferred_term,
                f"Class {cls!r}: YAML C-code {yaml_ccode} != CT C-code "
                f"{ct_entity.ccode} (CT row {ct_entity.row_number})")
        for field_name, yaml_key, rule in (
            ("Preferred Term", "Preferred Term", "CLASS_TERM_MISMATCH"),
            ("Definition", "Definition", "CLASS_DEFINITION_MISMATCH"),
        ):
            yaml_val = body.get(yaml_key)
            ct_val = ct_entity.preferred_term if rule.endswith("TERM_MISMATCH") else ct_entity.definition
            if yaml_val and ct_val and yaml_val.strip() != ct_val.strip():
                add(rule, "error", cls, yaml_ccode, ct_entity.ccode, yaml_val, ct_val,
                    f"Class {cls!r}: {field_name} differs between YAML and CT "
                    f"(CT row {ct_entity.row_number})")

    # --- attribute level ---------------------------------------------------
    for cls, body in model.items():
        if not isinstance(body, dict):
            continue
        for attr, abody in (body.get("Attributes") or {}).items():
            if not isinstance(abody, dict):
                continue
            location = f"{cls}.{attr}"
            yaml_ccode = abody.get("NCI C-Code")
            ct_row = ct.attributes.get((cls, attr))

            if yaml_ccode and not ct.code_owners(yaml_ccode):
                add("CCODE_NOT_IN_CT", "error", location, yaml_ccode, None,
                    abody.get("Preferred Term"), None,
                    f"{location}: YAML C-code {yaml_ccode} does not occur in USDM_CT.xlsx")

            if yaml_ccode and ct_row is None:
                add("ATTR_NOT_IN_CT", "warning", location, yaml_ccode, None,
                    abody.get("Preferred Term"), None,
                    f"{location}: YAML attribute carries C-code {yaml_ccode} but has no "
                    f"row in USDM_CT.xlsx")
                continue

            if ct_row is None or not yaml_ccode:
                continue

            if yaml_ccode != ct_row.ccode:
                owners = [o for o in ct.code_owners(yaml_ccode) if o != location]
                owner_note = (f"; YAML code is the CT row of {', '.join(owners)}"
                              if owners else "; YAML code not found elsewhere in CT")
                add("ATTR_CCODE_MISMATCH", "error", location, yaml_ccode, ct_row.ccode,
                    abody.get("Preferred Term"), ct_row.preferred_term,
                    f"{location}: YAML C-code {yaml_ccode} != CT C-code {ct_row.ccode} "
                    f"(CT row {ct_row.row_number}){owner_note}")

            yaml_term = abody.get("Preferred Term")
            if yaml_term and ct_row.preferred_term and yaml_term.strip() != ct_row.preferred_term.strip():
                add("ATTR_TERM_MISMATCH", "error", location, yaml_ccode, ct_row.ccode,
                    yaml_term, ct_row.preferred_term,
                    f"{location}: Preferred Term differs "
                    f"(YAML: {yaml_term!r}, CT row {ct_row.row_number}: {ct_row.preferred_term!r})")

            yaml_def = abody.get("Definition")
            if yaml_def and ct_row.definition and yaml_def.strip() != ct_row.definition.strip():
                add("ATTR_DEFINITION_MISMATCH", "error", location, yaml_ccode, ct_row.ccode,
                    yaml_def, ct_row.definition,
                    f"{location}: Definition differs between YAML and CT row {ct_row.row_number}")

            # value list vs YAML type consistency
            has_vl = bool(ct_row.has_value_list) and ct_row.has_value_list.startswith("Y")
            is_code_typed = bool(_type_refs(abody) & CODE_TYPES)
            if has_vl and not is_code_typed:
                add("VALUE_LIST_TYPE_MISMATCH", "warning", location, yaml_ccode, ct_row.ccode,
                    ", ".join(sorted(_type_refs(abody))), ct_row.has_value_list,
                    f"{location}: CT declares a value list ({ct_row.has_value_list}) but the "
                    f"YAML attribute is not typed as Code/AliasCode")
            elif not has_vl and is_code_typed and ct_row.has_value_list == "N":
                add("VALUE_LIST_TYPE_MISMATCH", "warning", location, yaml_ccode, ct_row.ccode,
                    ", ".join(sorted(_type_refs(abody))), ct_row.has_value_list,
                    f"{location}: YAML attribute is Code-typed but CT declares no value list "
                    f"(Has Value List = N)")

    # --- reverse coverage: CT rows without a YAML C-code counterpart --------
    for (entity, attr), ct_row in sorted(ct.attributes.items()):
        body = model.get(entity)
        if not isinstance(body, dict):
            continue  # class-level finding already reported
        abody = (body.get("Attributes") or {}).get(attr)
        if abody is None or not abody.get("NCI C-Code"):
            state = "missing from YAML" if abody is None else "has no C-code in YAML"
            add("CT_ATTR_NOT_IN_YAML", "info", f"{entity}.{attr}", None, ct_row.ccode,
                None, ct_row.preferred_term,
                f"{entity}.{attr}: CT row {ct_row.row_number} ({ct_row.role}, "
                f"{ct_row.ccode}) {state}")

    # --- classification definitions without a value list (issue #717) -------
    for (entity, attr), ct_row in sorted(ct.attributes.items()):
        if (ct_row.definition or "").startswith(CLASSIFICATION_PREFIX) and ct_row.has_value_list == "N":
            add("CLASSIFICATION_WITHOUT_VALUE_LIST", "warning", f"{entity}.{attr}",
                None, ct_row.ccode, None, ct_row.definition,
                f"{entity}.{attr}: definition starts with {CLASSIFICATION_PREFIX!r} "
                f"but Has Value List = N (CT row {ct_row.row_number})")

    return findings
